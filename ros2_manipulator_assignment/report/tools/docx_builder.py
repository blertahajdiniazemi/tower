"""Small helper layer over python-docx for an academic report.

Provides: A4 page set-up, styles, automatic multilevel heading numbering,
a real TOC field, SEQ-numbered captions with bookmarks, REF cross-references,
PAGE/NUMPAGES footer fields, tables with repeated header rows and code blocks.
"""
import copy
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = 'Times New Roman'
MONO_FONT = 'Courier New'
ACCENT = RGBColor(0x1F, 0x3B, 0x5A)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def _set_fonts(rpr, name):
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = _el('w:rFonts')
        rpr.insert(0, fonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(a), name)
    # theme font attributes would override the explicit font names
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]

PPR_ORDER = ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl',
             'numPr', 'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku',
             'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi',
             'adjustRightInd', 'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents',
             'suppressOverlap', 'jc', 'textDirection', 'textAlignment', 'textboxTightWrap',
             'outlineLvl', 'divId', 'cnfStyle', 'rPr', 'sectPr', 'pPrChange']
TBLPR_ORDER = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
               'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd',
               'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription']
TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap',
              'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark']
TRPR_ORDER = ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit',
              'trHeight', 'tblHeader', 'tblCellSpacing', 'jc', 'hidden']


def insert_ordered(parent, child, order):
    """Insert child into parent respecting the schema sequence given in order."""
    name = child.tag.split('}')[1]
    pos = order.index(name)
    for i, existing in enumerate(list(parent)):
        ename = existing.tag.split('}')[1]
        if ename in order and order.index(ename) > pos:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


class Report:
    def __init__(self, known_refs=None, toc_pages=None):
        self.doc = Document()
        self.known_refs = known_refs or {}   # refs collected on a previous pass (forward refs)
        self._bm_id = 100
        self.fig_no = 0
        self.tab_no = 0
        self.lst_no = 0
        self.refs = {}          # key -> (label, number)
        self.headings = []      # (level, number_text, title) for TOC cache
        self.toc_pages = toc_pages or {}     # title -> page (measured on a previous pass)
        self._h_counters = [0, 0, 0]
        self._setup_page()
        self._setup_styles()
        self._setup_heading_numbering()

    # ------------------------------------------------------------------ setup
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(2.5)
        s.header_distance = s.footer_distance = Cm(1.25)
        s.different_first_page_header_footer = True

    def _setup_styles(self):
        st = self.doc.styles
        # document defaults: replace theme fonts by the body font
        rpr_default = self.doc.styles.element.find(qn('w:docDefaults'))
        if rpr_default is not None:
            for fonts in rpr_default.iter(qn('w:rFonts')):
                _set_fonts(fonts.getparent(), BODY_FONT)
        normal = st['Normal']
        normal.font.name = BODY_FONT
        normal.font.size = Pt(12)
        _set_fonts(normal.element.get_or_add_rPr(), BODY_FONT)
        pf = normal.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.15
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        for name, size, before, after in (('Heading 1', 15, 18, 8), ('Heading 2', 13, 12, 6),
                                          ('Heading 3', 12, 10, 4)):
            h = st[name]
            h.font.name = BODY_FONT
            _set_fonts(h.element.get_or_add_rPr(), BODY_FONT)
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = ACCENT
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.line_spacing = 1.1
        cap = st['Caption']
        cap.font.name = BODY_FONT
        _set_fonts(cap.element.get_or_add_rPr(), BODY_FONT)
        cap.font.size = Pt(10)
        cap.font.italic = False
        cap.font.bold = False
        cap.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        cap.paragraph_format.space_before = Pt(3)
        cap.paragraph_format.space_after = Pt(10)
        cap.paragraph_format.line_spacing = 1.0
        # code block style
        from docx.enum.style import WD_STYLE_TYPE
        code = st.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        code.font.name = MONO_FONT
        _set_fonts(code.element.get_or_add_rPr(), MONO_FONT)
        code.font.size = Pt(9)
        cpf = code.paragraph_format
        cpf.line_spacing = 1.0
        cpf.space_after = Pt(0)
        cpf.space_before = Pt(0)
        # hanging indent: a long line that wraps continues visibly indented under its start
        cpf.left_indent = Cm(1.2)
        cpf.first_line_indent = Cm(-1.0)
        ppr = code.element.get_or_add_pPr()
        shd = _el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': 'F3F5F7'})
        insert_ordered(ppr, shd, PPR_ORDER)
        tstyle = st.add_style('Table Text', WD_STYLE_TYPE.PARAGRAPH)
        tstyle.base_style = normal
        tstyle.font.size = Pt(9.5)
        tstyle.paragraph_format.space_after = Pt(0)
        tstyle.paragraph_format.line_spacing = 1.05
        for lvl in (1, 2):
            try:
                t = st['TOC %d' % lvl]
            except KeyError:
                t = st.add_style('TOC %d' % lvl, WD_STYLE_TYPE.PARAGRAPH)
                t.base_style = normal
            t.paragraph_format.space_after = Pt(1 if lvl == 2 else 3)
            t.paragraph_format.space_before = Pt(0)
            t.paragraph_format.line_spacing = 1.0
            t.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
            t.font.size = Pt(11 if lvl == 1 else 10.5)
            t.font.bold = lvl == 1
        lb = st['List Bullet']
        lb.paragraph_format.space_after = Pt(3)

    def _setup_heading_numbering(self):
        numbering = self.doc.part.numbering_part.element
        abs_id = 90
        absn = _el('w:abstractNum', **{'w:abstractNumId': abs_id})
        absn.append(_el('w:multiLevelType', **{'w:val': 'multilevel'}))
        for ilvl, (fmt, style) in enumerate((('%1.', 'Heading1'), ('%1.%2', 'Heading2'),
                                             ('%1.%2.%3', 'Heading3'))):
            lvl = _el('w:lvl', **{'w:ilvl': ilvl})
            lvl.append(_el('w:start', **{'w:val': 1}))
            lvl.append(_el('w:numFmt', **{'w:val': 'decimal'}))
            lvl.append(_el('w:pStyle', **{'w:val': style}))
            lvl.append(_el('w:lvlText', **{'w:val': fmt}))
            lvl.append(_el('w:lvlJc', **{'w:val': 'left'}))
            ppr = _el('w:pPr')
            ind = {0: 567, 1: 680, 2: 850}[ilvl]
            ppr.append(_el('w:ind', **{'w:left': ind, 'w:hanging': ind}))
            lvl.append(ppr)
            absn.append(lvl)
        # abstractNum elements must precede num elements
        first_num = numbering.find(qn('w:num'))
        if first_num is not None:
            first_num.addprevious(absn)
        else:
            numbering.append(absn)
        num = _el('w:num', **{'w:numId': 90})
        num.append(_el('w:abstractNumId', **{'w:val': abs_id}))
        numbering.append(num)
        for ilvl, name in enumerate(('Heading 1', 'Heading 2', 'Heading 3')):
            ppr = self.doc.styles[name].element.get_or_add_pPr()
            numpr = _el('w:numPr')
            numpr.append(_el('w:ilvl', **{'w:val': ilvl}))
            numpr.append(_el('w:numId', **{'w:val': 90}))
            insert_ordered(ppr, numpr, PPR_ORDER)

    # ------------------------------------------------------------- primitives
    def _bookmark(self, paragraph, name, runs_start_index=0):
        self._bm_id += 1
        start = _el('w:bookmarkStart', **{'w:id': self._bm_id, 'w:name': name})
        end = _el('w:bookmarkEnd', **{'w:id': self._bm_id})
        return start, end

    def _field(self, paragraph, instr, cached, bold=False, italic=False, size=None):
        """Append a complex field (begin/instr/separate/cached result/end)."""
        def run_with(child, text=None):
            r = _el('w:r')
            rpr = _el('w:rPr')
            if bold:
                rpr.append(_el('w:b'))
            if italic:
                rpr.append(_el('w:i'))
            if size:
                rpr.append(_el('w:sz', **{'w:val': int(size * 2)}))
            if len(rpr):
                r.append(rpr)
            r.append(child)
            if text is not None:
                child.text = text
                child.set(qn('xml:space'), 'preserve')
            paragraph._p.append(r)
            return r
        run_with(_el('w:fldChar', **{'w:fldCharType': 'begin'}))
        run_with(_el('w:instrText'), ' %s ' % instr)
        run_with(_el('w:fldChar', **{'w:fldCharType': 'separate'}))
        run_with(_el('w:t'), cached)
        run_with(_el('w:fldChar', **{'w:fldCharType': 'end'}))

    def inline(self, paragraph, text, size=None, color=None):
        """Add text with light markup: **bold**, *italic*, `code`, {ref:key}, [n] citations."""
        tokens = re.split(r'(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|\{ref:[A-Za-z0-9_\-]+\})', text)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith('{ref:'):
                self.ref(paragraph, tok[5:-1])
                continue
            if tok.startswith('**') and tok.endswith('**'):
                r = paragraph.add_run(tok[2:-2])
                r.bold = True
            elif tok.startswith('`') and tok.endswith('`'):
                r = paragraph.add_run(tok[1:-1])
                r.font.name = MONO_FONT
                _set_fonts(r._r.get_or_add_rPr(), MONO_FONT)
                r.font.size = Pt((size or 12) - 1.5)
            elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
                r = paragraph.add_run(tok[1:-1])
                r.italic = True
            else:
                r = paragraph.add_run(tok)
            if size and not tok.startswith('`'):
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return paragraph

    # --------------------------------------------------------------- content
    def heading(self, text, level=1, numbered=True, page_break=False):
        if page_break:
            self.page_break()
        p = self.doc.add_paragraph(style='Heading %d' % level)
        if numbered:
            self._h_counters[level - 1] += 1
            for i in range(level, 3):
                self._h_counters[i] = 0
            number = '.'.join(str(c) for c in self._h_counters[:level])
            number = number + '.' if level == 1 else number
        else:
            ppr = p._p.get_or_add_pPr()
            numpr = _el('w:numPr')
            numpr.append(_el('w:ilvl', **{'w:val': 0}))
            numpr.append(_el('w:numId', **{'w:val': 0}))
            insert_ordered(ppr, numpr, PPR_ORDER)
            number = ''
        # bookmark for TOC hyperlinks
        self._bm_id += 1
        name = '_Toc%08d' % self._bm_id
        p._p.append(_el('w:bookmarkStart', **{'w:id': self._bm_id, 'w:name': name}))
        p.add_run(text)
        p._p.append(_el('w:bookmarkEnd', **{'w:id': self._bm_id}))
        if level <= 2:
            self.headings.append((level, number, text, name))
        return p

    def para(self, text, style=None, align=None, size=None, space_after=None, keep_next=False,
             italic=False):
        p = self.doc.add_paragraph(style=style)
        self.inline(p, text, size=size)
        if italic:
            for r in p.runs:
                r.italic = True
        if align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == 'justify':
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == 'right':
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if keep_next:
            p.paragraph_format.keep_with_next = True
        return p

    def body(self, text):
        return self.para(text, align='justify')

    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph(style='List Bullet')
            self.inline(p, it)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        p.paragraph_format.space_after = Pt(0)

    def _caption(self, label, key, text, attribution=None, container=None):
        if label == 'Figure':
            self.fig_no += 1
            n = self.fig_no
        elif label == 'Table':
            self.tab_no += 1
            n = self.tab_no
        else:
            self.lst_no += 1
            n = self.lst_no
        self.refs[key] = (label, n)
        p = (container or self.doc).add_paragraph(style='Caption')
        self._bm_id += 1
        bm = '_Ref_%s' % key.replace('-', '_')
        p._p.append(_el('w:bookmarkStart', **{'w:id': self._bm_id, 'w:name': bm}))
        r = p.add_run(label + ' ')
        r.bold = True
        self._field(p, 'SEQ %s \\* ARABIC' % label, str(n), bold=True)
        p._p.append(_el('w:bookmarkEnd', **{'w:id': self._bm_id}))
        r = p.add_run('. ')
        r.bold = True
        self.inline(p, text, size=10)
        if attribution:
            self.inline(p, ' ' + attribution, size=10, color=RGBColor(0x55, 0x55, 0x55))
        return p

    def ref(self, paragraph, key):
        label, n = self.refs.get(key) or self.known_refs.get(key) or ('??', 0)
        self._field(paragraph, 'REF _Ref_%s \\h' % key.replace('-', '_'), '%s %d' % (label, n))

    def figure(self, path, key, caption, width_cm=16.0, attribution=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Cm(width_cm))
        cap = self._caption('Figure', key, caption, attribution)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return cap

    def figure_pair(self, left, right, gap_cm=0.6):
        """Two figures side by side; each item is (path, key, caption, width_cm, attribution)."""
        total = 16.0
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        insert_ordered(tblpr, _el('w:tblLayout', **{'w:type': 'fixed'}), TBLPR_ORDER)
        borders = _el('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(_el('w:' + side, **{'w:val': 'nil'}))
        insert_ordered(tblpr, borders, TBLPR_ORDER)
        widths = [left[3] + gap_cm / 2, right[3] + gap_cm / 2]
        scale = total / sum(widths)
        widths = [w * scale for w in widths]
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(int(widths[i] / 2.54 * 1440)))
        insert_ordered(t.rows[0]._tr.get_or_add_trPr(), _el('w:cantSplit'), TRPR_ORDER)
        for cell, item, w in zip(t.rows[0].cells, (left, right), widths):
            path, key, caption, width_cm, attribution = item
            cell.width = Cm(w)
            tcpr = cell._tc.get_or_add_tcPr()
            insert_ordered(tcpr, _el('w:vAlign', **{'w:val': 'bottom'}), TCPR_ORDER)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(path, width=Cm(width_cm))
            cap = self._caption('Figure', key, caption, attribution, container=cell)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        return t

    def table(self, header, rows, key, caption, widths_cm, font_size=9.5, attribution=None,
              caption_above=True, bold_first_col=False):
        if caption_above:
            cap = self._caption('Table', key, caption, attribution)
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_after = Pt(4)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        insert_ordered(tblpr, _el('w:tblLayout', **{'w:type': 'fixed'}), TBLPR_ORDER)
        # header row
        hdr = t.rows[0]
        trpr = hdr._tr.get_or_add_trPr()
        insert_ordered(trpr, _el('w:cantSplit'), TRPR_ORDER)
        insert_ordered(trpr, _el('w:tblHeader'), TRPR_ORDER)
        for i, h in enumerate(header):
            c = hdr.cells[i]
            c.width = Cm(widths_cm[i])
            self._cell(c, h, font_size, bold=True, fill='DCE6F0')
        for row in rows:
            r = t.add_row()
            insert_ordered(r._tr.get_or_add_trPr(), _el('w:cantSplit'), TRPR_ORDER)
            for i, val in enumerate(row):
                c = r.cells[i]
                c.width = Cm(widths_cm[i])
                self._cell(c, val, font_size, bold=(bold_first_col and i == 0))
        if len(rows) <= 4:
            # short tables are kept on one page: every row except the last keeps with the next
            for tr in t.rows[:-1]:
                for c in tr.cells:
                    for p in c.paragraphs:
                        p.paragraph_format.keep_with_next = True
        # grid widths
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(int(widths_cm[i] / 2.54 * 1440)))
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 0.8
        if not caption_above:
            self._caption('Table', key, caption, attribution)
        return t

    def _cell(self, cell, text, size, bold=False, fill=None):
        cell.paragraphs[0].style = self.doc.styles['Table Text']
        parts = text.split('\n')
        for j, part in enumerate(parts):
            p = cell.paragraphs[0] if j == 0 else cell.add_paragraph(style='Table Text')
            self.inline(p, part, size=size)
            if bold:
                for r in p.runs:
                    r.bold = True
        tcpr = cell._tc.get_or_add_tcPr()
        mar = _el('w:tcMar')
        for side, width in (('top', 40), ('left', 80), ('bottom', 40), ('right', 80)):
            mar.append(_el('w:' + side, **{'w:w': width, 'w:type': 'dxa'}))
        insert_ordered(tcpr, mar, TCPR_ORDER)
        if fill:
            insert_ordered(tcpr, _el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': fill}),
                           TCPR_ORDER)

    def code(self, text, key, caption, attribution=None):
        cap = self._caption('Listing', key, caption, attribution)
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_after = Pt(3)
        lines = text.rstrip('\n').split('\n')
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph(style='Code Block')
            r = p.add_run(line if line else ' ')
            r.font.name = MONO_FONT
            if i < len(lines) - 1:
                p.paragraph_format.keep_with_next = True
            if i == 0:
                p.paragraph_format.space_before = Pt(2)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    # ------------------------------------------------------ front/back matter
    def toc(self, title='Table of Contents'):
        p = self.doc.add_paragraph(style='Heading 1')
        ppr = p._p.get_or_add_pPr()
        numpr = _el('w:numPr')
        numpr.append(_el('w:ilvl', **{'w:val': 0}))
        numpr.append(_el('w:numId', **{'w:val': 0}))
        insert_ordered(ppr, numpr, PPR_ORDER)
        p.add_run(title)
        self._toc_anchor = self.doc.add_paragraph()   # placeholder, filled by finalize_toc()

    def finalize_toc(self):
        """Replace the placeholder with a TOC field whose cached result lists the headings
        with hyperlinks and the page numbers measured on the previous rendering pass."""
        anchor = self._toc_anchor._p
        body = anchor.getparent()
        idx = list(body).index(anchor)
        paras = []

        def mk_par(style):
            p = _el('w:p')
            ppr = _el('w:pPr')
            ppr.append(_el('w:pStyle', **{'w:val': style.replace(' ', '')}))
            tabs = _el('w:tabs')
            tabs.append(_el('w:tab', **{'w:val': 'right', 'w:leader': 'dot', 'w:pos': 9060}))
            ppr.append(tabs)
            p.append(ppr)
            return p

        def run(parent, child, text=None):
            r = _el('w:r')
            r.append(child)
            if text is not None:
                child.text = text
                child.set(qn('xml:space'), 'preserve')
            parent.append(r)
        entries = self.headings
        for i, (level, number, title, bm) in enumerate(entries):
            p = mk_par('TOC %d' % level)
            if i == 0:
                run(p, _el('w:fldChar', **{'w:fldCharType': 'begin'}))
                run(p, _el('w:instrText'), ' TOC \\o "1-2" \\h \\z \\u ')
                run(p, _el('w:fldChar', **{'w:fldCharType': 'separate'}))
            h = _el('w:hyperlink', **{'w:anchor': bm, 'w:history': 1})
            label = ('%s  %s' % (number, title)) if number else title
            run(h, _el('w:t'), label)
            run(h, _el('w:tab'))
            page = self.toc_pages.get(title, '')
            run(h, _el('w:t'), str(page))
            p.append(h)
            if i == len(entries) - 1:
                run(p, _el('w:fldChar', **{'w:fldCharType': 'end'}))
            paras.append(p)
        body.remove(anchor)
        for k, p in enumerate(paras):
            body.insert(idx + k, p)

    def header_footer(self, header_text, footer_left):
        sec = self.doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.text = ''
        r = hp.add_run(header_text)
        r.font.size = Pt(9)
        r.italic = True
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pbdr = _el('w:pBdr')
        pbdr.append(_el('w:bottom', **{'w:val': 'single', 'w:sz': 4, 'w:space': 1, 'w:color': 'A0A0A0'}))
        insert_ordered(hp._p.get_or_add_pPr(), pbdr, PPR_ORDER)
        fp = sec.footer.paragraphs[0]
        fp.text = ''
        ppr = fp._p.get_or_add_pPr()
        tabs = _el('w:tabs')
        # clear the Footer style's default centre/right tabs, then right-align at the text margin
        tabs.append(_el('w:tab', **{'w:val': 'clear', 'w:pos': 4680}))
        tabs.append(_el('w:tab', **{'w:val': 'right', 'w:pos': 9070}))
        tabs.append(_el('w:tab', **{'w:val': 'clear', 'w:pos': 9360}))
        insert_ordered(ppr, tabs, PPR_ORDER)
        r = fp.add_run(footer_left + '\t')
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        r = fp.add_run('Page ')
        r.font.size = Pt(9)
        self._field(fp, 'PAGE', '1', size=9)
        r = fp.add_run(' of ')
        r.font.size = Pt(9)
        self._field(fp, 'NUMPAGES', '1', size=9)

    def update_fields_on_open(self):
        settings = self.doc.settings.element
        later = ('hdrShapeDefaults', 'footnotePr', 'endnotePr', 'compat', 'docVars', 'rsids',
                 'mathPr', 'attachedSchema', 'themeFontLang', 'clrSchemeMapping',
                 'doNotIncludeSubdocsInStats', 'doNotAutoCompressPictures', 'forceUpgrade',
                 'captions', 'readModeInkLockDown', 'smartTagType', 'schemaLibrary',
                 'shapeDefaults', 'doNotEmbedSmartTags', 'decimalSymbol', 'listSeparator')
        uf = _el('w:updateFields', **{'w:val': 'true'})
        for child in settings:
            if child.tag.split('}')[1] in later:
                child.addprevious(uf)
                break
        else:
            settings.append(uf)
        zoom = settings.find(qn('w:zoom'))
        if zoom is not None and zoom.get(qn('w:percent')) is None:
            zoom.set(qn('w:percent'), '100')

    def save(self, path):
        self.doc.save(path)
